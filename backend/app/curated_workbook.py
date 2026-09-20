"""只读解析负责人确认的公司资料表；不计算公式，不访问链接。"""

from __future__ import annotations

import hashlib
import io
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Literal
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from backend.app.providers import _validate_public_http_url, validate_unified_credit_code

PARSER_VERSION = "curated-xlsx-v1"
PRIVATE_ROOT = Path(__file__).resolve().parents[2] / "data/private/research_imports"
ImportMode = Literal["identity_only", "initial_data"]
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_ROWS = 1000
HEADERS = {
    "公司总表": [
        "公司ID",
        "项目简称",
        "工商全称",
        "统一社会信用代码",
        "地区（资料口径）",
        "主体来源1",
        "主体公开链接1",
        "主体来源2",
        "主体公开链接2",
    ],
    "事件明细": [
        "事件ID",
        "公司ID",
        "事件/披露日期",
        "事件类别",
        "重要事件",
        "事件摘要",
        "证据等级",
        "内容支持状态",
        "工商全称",
        "统一社会信用代码",
        "日期精度",
        "日期口径",
        "实际发生日期",
        "披露日期",
        "主体归属口径",
        "交易/进程组",
        "信息来源1",
        "公开链接1",
        "信息来源2",
        "公开链接2",
        "核验说明",
        "资料截止日期",
    ],
    "证据来源": [
        "来源ID",
        "来源属性",
        "发布者/载体",
        "标题",
        "公开URL",
        "可读取范围",
        "证据定位提示",
        "来源使用限制",
    ],
    "待核事项": [
        "线索ID",
        "公司ID",
        "资料日期",
        "待核主题",
        "目前信息",
        "处理口径",
        "公开链接/定位",
        "补证要求",
    ],
}
CATEGORY_MAP = {
    "融资": "financing_cap_table",
    "股权/股东变化": "financing_cap_table",
    "IPO进度": "exit_liquidity",
    "经营与财务": "financial_operation",
    "经营财务": "financial_operation",
    "融资与股权": "financing_cap_table",
    "合同与商业进展": "contract_commercial",
    "合同商业": "contract_commercial",
    "产品与技术": "product_technology",
    "产品技术": "product_technology",
    "治理与人员": "governance_people",
    "治理人员": "governance_people",
    "司法与合规": "legal_compliance",
    "司法合规": "legal_compliance",
    "产能与资产": "capacity_assets",
    "产能资产": "capacity_assets",
    "退出进程": "exit_liquidity",
}
CATEGORY_MAP.update({value: value for value in tuple(CATEGORY_MAP.values())})


def digest(value: object) -> str:
    return hashlib.sha256(
        json.dumps(
            value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str
        ).encode()
    ).hexdigest()


def _text(value: object, *, required: bool = False, limit: int = 2000) -> str:
    if isinstance(value, datetime):
        value = value.date().isoformat()
    elif isinstance(value, date):
        value = value.isoformat()
    if value is None:
        value = ""
    if not isinstance(value, str):
        raise ValueError("文本字段类型不符")
    value = value.strip()
    if value.startswith("="):
        raise ValueError("资料字段不接受公式；请使用已确认的文本值")
    if len(value) > limit or (required and not value):
        raise ValueError("资料字段为空或超过长度上限")
    return value


def exact_day(value: str) -> date | None:
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return date.fromisoformat(value)
    return None


@dataclass(frozen=True)
class CuratedSource:
    url: str
    name: str
    metadata: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class CuratedCompany:
    key: str
    row: int
    legal_name: str
    credit_code: str
    alias: str
    region: str
    sources: tuple[CuratedSource, ...]

    def identity(self) -> dict[str, object]:
        # 不携带汇总列、业务日期或自由文本归集说明进入独立检索的身份输入。
        return {
            "legal_name": self.legal_name,
            "credit_code": self.credit_code,
            "alias": self.alias,
            "region": self.region,
            "sources": [{"url": s.url, "name": s.name} for s in self.sources],
        }


@dataclass(frozen=True)
class CuratedRecord:
    key: str
    company_key: str
    sheet: str
    row: int
    title: str
    summary: str
    event_type: str
    values: dict[str, str]
    sources: tuple[CuratedSource, ...]
    confirmed: bool

    def public_metadata(self) -> dict[str, object]:
        return {
            "date_text": self.values.get("事件/披露日期", ""),
            "date_precision": self.values.get("日期精度", "未知"),
            "date_basis": self.values.get("日期口径", "资料日期"),
            "occurred_date_text": self.values.get("实际发生日期", "未知"),
            "as_of_date": self.values.get("资料截止日期", ""),
            "subject_scope": self.values.get("主体归属口径", ""),
            "source_grade": self.values.get("证据等级", "待核线索"),
            "content_support": self.values.get("内容支持状态", "待核"),
            "assessment_status": "not_assessed",
        }

    def event_fields(self) -> dict[str, object]:
        facts = [
            {"name": "重要事件", "value": self.title, "unit": None},
            {"name": "人工整理摘要", "value": self.summary, "unit": None},
        ]
        for name in ("主体归属口径", "事件/披露日期", "日期精度", "日期口径", "实际发生日期"):
            if self.values.get(name):
                facts.append({"name": name, "value": self.values[name], "unit": None})
        notes = [
            self.values[key] for key in ("核验说明", "补证要求", "处理口径") if self.values.get(key)
        ]
        publication_day = exact_day(self.values.get("披露日期", ""))
        return {
            "title": self.title,
            "summary": self.summary,
            "facts": facts,
            "uncertainties": notes,
            "event_type": self.event_type,
            "event_subtype": "curated_record",
            "occurred_at": None,
            "published_at": None,
            "published_on": publication_day.isoformat() if publication_day else None,
        }


@dataclass(frozen=True)
class WorkbookIssue:
    sheet: str
    row: int
    company_key: str
    message: str


@dataclass(frozen=True)
class LoadedCuratedWorkbook:
    file_hash: str
    source_filename: str
    dataset_key: str
    mode: ImportMode
    company_keys: tuple[str, ...]
    companies: tuple[CuratedCompany, ...]
    records: tuple[CuratedRecord, ...]
    issues: tuple[WorkbookIssue, ...]
    notes: tuple[dict[str, object], ...] = ()

    @property
    def selection_key(self) -> str:
        # 同一原文件可先导入选中记录再补齐全表；历史回执保留，事项指纹仍稳定。
        return digest(
            [
                "curated-selection-v2",
                self.dataset_key,
                self.mode,
                self.company_keys,
                sorted((r.sheet, r.key, digest(asdict(r))) for r in self.records),
                self.notes,
            ]
        )


def _rows(workbook, sheet: str):
    if sheet not in workbook.sheetnames:
        raise ValueError(f"缺少工作表：{sheet}")
    iterator = workbook[sheet].iter_rows(max_col=32, values_only=True)
    row_number = 0
    for row_number, cells in enumerate(iterator, 1):
        if row_number < 5:
            continue
        if row_number == 5:
            names = [str(v).strip() if v is not None else "" for v in cells]
            if any(names.count(name) != 1 for name in HEADERS[sheet]):
                raise ValueError(f"{sheet} 第 5 行列名缺失或重复")
            indexes = {name: names.index(name) for name in HEADERS[sheet]}
            continue
        if row_number > MAX_ROWS:
            raise ValueError(f"{sheet} 超过 {MAX_ROWS} 行上限")
        if any(cells[i] is not None for i in indexes.values()):
            yield row_number, {name: cells[i] for name, i in indexes.items()}
    if row_number < 5:
        raise ValueError(f"{sheet} 缺少表头")


def _sources(values: dict[str, str], *, identity: bool = False) -> tuple[CuratedSource, ...]:
    found = {}
    for index in (1, 2):
        url = values.get(f"主体公开链接{index}" if identity else f"公开链接{index}", "")
        if not url:
            continue
        url = _validate_public_http_url(url)
        name = values.get(f"主体来源{index}" if identity else f"信息来源{index}", "")
        found[url] = CuratedSource(url, name or "人工整理来源")
    return tuple(found.values())


class CuratedWorkbookProvider:
    code = "curated_workbook_import"
    parser_version = PARSER_VERSION
    external_calls = 0

    def __init__(
        self,
        file_path: str | Path,
        *,
        dataset_key: str,
        mode: ImportMode,
        company_keys: tuple[str, ...] = (),
        allowed_root: Path | None = None,
    ):
        if mode not in {"identity_only", "initial_data"}:
            raise ValueError("invalid curated import mode")
        if not re.fullmatch(r"[a-zA-Z0-9_-]{1,64}", dataset_key):
            raise ValueError("dataset_key must be a stable ASCII identifier")
        self.path = Path(file_path)
        self.root = (allowed_root or PRIVATE_ROOT).resolve()
        self.dataset_key, self.mode = dataset_key, mode
        self.company_keys = tuple(sorted(set(company_keys)))

    def load(self) -> LoadedCuratedWorkbook:
        path = self.path if self.path.is_absolute() else self.root / self.path
        path = path.resolve()
        if not path.is_relative_to(self.root) or path.suffix.lower() != ".xlsx":
            raise ValueError("curated workbook must be an XLSX inside the private import directory")
        with path.open("rb") as stream:
            content = stream.read(MAX_FILE_BYTES + 1)
        if len(content) > MAX_FILE_BYTES:
            raise ValueError("workbook exceeds 5 MiB limit")
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                files = archive.infolist()
                if len(files) > 200 or sum(item.file_size for item in files) > MAX_EXPANDED_BYTES:
                    raise ValueError("workbook expanded size exceeds limit")
                if any("vbaproject" in item.filename.lower() for item in files):
                    raise ValueError("macro-enabled workbooks are not accepted")
        except BadZipFile as error:
            raise ValueError("invalid XLSX archive") from error
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        issues, companies, records, notes = [], [], [], []
        try:
            for row, raw in _rows(workbook, "公司总表"):
                key = str(raw["公司ID"] or "").strip()
                try:
                    values = {k: _text(v) for k, v in raw.items()}
                    company = CuratedCompany(
                        key=_text(values["公司ID"], required=True, limit=64),
                        row=row,
                        legal_name=_text(values["工商全称"], required=True, limit=240),
                        credit_code=validate_unified_credit_code(values["统一社会信用代码"]),
                        alias=_text(values["项目简称"], limit=240),
                        region=_text(values["地区（资料口径）"], limit=120),
                        sources=_sources(values, identity=True),
                    )
                    companies.append(company)
                except ValueError as error:
                    issues.append(WorkbookIssue("公司总表", row, key, str(error)))
            duplicate_keys = {
                c.key for c in companies if sum(o.key == c.key for o in companies) > 1
            }
            for company in companies:
                if company.key in duplicate_keys or any(
                    other.key != company.key
                    and (
                        other.credit_code == company.credit_code
                        or other.legal_name == company.legal_name
                    )
                    for other in companies
                ):
                    issues.append(
                        WorkbookIssue(
                            "公司总表",
                            company.row,
                            company.key,
                            "公司引用、信用代码或名称重复；需要定点核对",
                        )
                    )
            selected = [c for c in companies if not self.company_keys or c.key in self.company_keys]
            known_keys = {c.key for c in companies}
            for key in set(self.company_keys) - known_keys:
                issues.append(WorkbookIssue("公司总表", 0, key, "未找到有效公司记录"))
            if self.mode == "initial_data":
                by_key = {c.key: c for c in selected}
                for sheet in ("事件明细", "待核事项"):
                    for row, raw in _rows(workbook, sheet):
                        key = str(raw["公司ID"] or "").strip()
                        if sheet == "待核事项" and (key == "多家" or "/" in key):
                            references = tuple(key.split("/")) if key != "多家" else ()
                            if (
                                self.company_keys
                                and references
                                and not set(references).intersection(self.company_keys)
                            ):
                                continue
                            try:
                                if any(reference not in known_keys for reference in references):
                                    raise ValueError("跨公司说明包含未知公司引用")
                                values = {k: _text(v) for k, v in raw.items()}
                                notes.append(
                                    {
                                        "sheet": sheet,
                                        "row": row,
                                        "company_keys": list(references),
                                        "values": values,
                                        "kind": "dataset_note"
                                        if not references
                                        else "company_note",
                                    }
                                )
                            except ValueError as error:
                                issues.append(WorkbookIssue(sheet, row, key, str(error)))
                            continue
                        if self.company_keys and key not in self.company_keys:
                            continue
                        try:
                            if key not in by_key:
                                raise ValueError("事项没有有效的公司引用")
                            values = {k: _text(v) for k, v in raw.items()}
                            if sheet == "事件明细":
                                company = by_key[key]
                                if (
                                    values["工商全称"] != company.legal_name
                                    or values["统一社会信用代码"] != company.credit_code
                                ):
                                    raise ValueError("事项主体与公司总表不一致")
                                category = CATEGORY_MAP.get(values["事件类别"])
                                if category is None:
                                    raise ValueError("本次表格适配不支持该事件类别")
                                sources = _sources(values)
                                support = values["内容支持状态"]
                                confirmed = bool(sources) and (
                                    support.startswith("已核对") or support == "媒体报道支持"
                                )
                                if any(
                                    word in values["重要事件"] + values["事件摘要"] + support
                                    for word in (
                                        "待核",
                                        "来源冲突",
                                        "欺诈",
                                        "刑事",
                                        "失联",
                                        "破产",
                                        "停业",
                                        "被调查",
                                    )
                                ):
                                    confirmed = False
                                title, summary, record_key = (
                                    values["重要事件"],
                                    values["事件摘要"],
                                    values["事件ID"],
                                )
                                for name in (
                                    "事件/披露日期",
                                    "披露日期",
                                    "实际发生日期",
                                    "资料截止日期",
                                ):
                                    exact_day(values[name])
                            else:
                                url = values["公开链接/定位"]
                                sources = (
                                    (CuratedSource(_validate_public_http_url(url), "待核来源"),)
                                    if url.startswith(("https://", "http://"))
                                    else ()
                                )
                                category, confirmed = "information_quality", False
                                title, summary, record_key = (
                                    values["待核主题"],
                                    values["目前信息"],
                                    values["线索ID"],
                                )
                                values["事件/披露日期"] = values["资料日期"]
                            record = CuratedRecord(
                                key=_text(record_key, required=True, limit=64),
                                company_key=key,
                                sheet=sheet,
                                row=row,
                                title=_text(title, required=True, limit=200),
                                summary=_text(summary, required=True, limit=950),
                                event_type=category,
                                values=values,
                                sources=sources,
                                confirmed=confirmed,
                            )
                            if any(
                                len(f"人工整理记录（非网页原文）\n{f['name']}：{f['value']}") > 1000
                                for f in record.event_fields()["facts"]
                            ):
                                raise ValueError("资料字段超过单条证据摘录上限")
                            records.append(record)
                        except ValueError as error:
                            issues.append(WorkbookIssue(sheet, row, key, str(error)))
                duplicate_ids = {
                    (r.sheet, r.key)
                    for r in records
                    if sum((o.sheet, o.key) == (r.sheet, r.key) for o in records) > 1
                }
                for record in records:
                    identifier = (record.sheet, record.key)
                    if identifier in duplicate_ids:
                        issues.append(
                            WorkbookIssue(
                                record.sheet, record.row, record.company_key, "事项记录 ID 重复"
                            )
                        )
                records = [r for r in records if (r.sheet, r.key) not in duplicate_ids]
                wanted_urls = {s.url for r in records for s in r.sources}
                metadata = {}
                for row, raw in _rows(workbook, "证据来源"):
                    if raw["公开URL"] not in wanted_urls:
                        continue
                    try:
                        values = {k: _text(v) for k, v in raw.items()}
                        metadata[values["公开URL"]] = values
                    except ValueError as error:
                        for key in {
                            r.company_key
                            for r in records
                            if any(s.url == raw["公开URL"] for s in r.sources)
                        }:
                            issues.append(WorkbookIssue("证据来源", row, key, str(error)))
                records = [
                    CuratedRecord(
                        **{
                            **r.__dict__,
                            "sources": tuple(
                                CuratedSource(s.url, s.name, metadata.get(s.url, {}))
                                for s in r.sources
                            ),
                        }
                    )
                    for r in records
                ]
            return LoadedCuratedWorkbook(
                hashlib.sha256(content).hexdigest(),
                path.name,
                self.dataset_key,
                self.mode,
                self.company_keys,
                tuple(selected),
                tuple(records),
                tuple(issues),
                tuple(notes),
            )
        finally:
            workbook.close()
